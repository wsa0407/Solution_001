from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from solution_gen.domain import DOC_TYPE_LABEL, Job
from solution_gen.store import output_dir
from solution_gen.workers.prose import chapter_export_text


def export_docx(job: Job, path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(12)
    if style.element.rPr is not None:
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    p = job.profile
    title = f"{p.place if p else ''}{p.industry if p else ''}{DOC_TYPE_LABEL.get(job.doc_type, '方案')}"
    h = doc.add_paragraph(title)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        run.bold = True
        run.font.size = Pt(18)
        run.font.name = "SimHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

    meta: list[str] = []
    if p and p.project_name:
        meta.append(p.project_name)
    if p and p.construction_unit:
        meta.append(p.construction_unit)
    if meta:
        sub = doc.add_paragraph(" · ".join(meta))
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, ch in enumerate(job.chapters, 1):
        doc.add_heading(f"{i}. {ch.title}", level=1)
        for para in chapter_export_text(ch.content, index=i, title=ch.title).split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        if ch.diagram:
            png = output_dir(job.id) / f"{ch.id}.png"
            if png.exists():
                doc.add_picture(str(png), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
